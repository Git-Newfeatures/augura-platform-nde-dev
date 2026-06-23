"""Server-side parsing of data files (CSV via stdlib, XLSX via openpyxl).

Returns a list of sheets {name, headers, rows} — raw data as str
(profiling and DQ are applied afterwards). No binary .xls (415).
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass

from openpyxl import load_workbook

from augura_api.core.errors import (
    BadRequestError,
    PayloadTooLargeError,
    UnsupportedMediaTypeError,
)

MAX_UPLOAD_BYTES = 25 * 1024 * 1024  # 25 MB


@dataclass(frozen=True)
class Sheet:
    name: str
    headers: list[str]
    rows: list[list[str]]


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise BadRequestError("unreadable file (unsupported encoding)")


def _parse_csv(data: bytes) -> Sheet:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = [[(c or "").strip() for c in row] for row in reader if any(c.strip() for c in row)]
    if not rows:
        raise BadRequestError("empty CSV")
    headers = rows[0]
    return Sheet(name="data", headers=headers, rows=rows[1:])


def _parse_xlsx(data: bytes) -> list[Sheet]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets: list[Sheet] = []
    for ws in wb.worksheets:
        rows_iter = ws.iter_rows(values_only=True)
        all_rows = [
            ["" if c is None else str(c) for c in row]
            for row in rows_iter
            if any(c is not None and str(c).strip() for c in row)
        ]
        if not all_rows:
            continue
        sheets.append(Sheet(name=ws.title, headers=all_rows[0], rows=all_rows[1:]))
    wb.close()
    if not sheets:
        raise BadRequestError("empty XLSX workbook")
    return sheets


def parse_upload(filename: str, data: bytes) -> list[Sheet]:
    if len(data) > MAX_UPLOAD_BYTES:
        raise PayloadTooLargeError("file too large", max_bytes=MAX_UPLOAD_BYTES, size=len(data))
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "csv":
        return [_parse_csv(data)]
    if ext == "xlsx":
        return _parse_xlsx(data)
    raise UnsupportedMediaTypeError("unsupported format (use .csv or .xlsx)", ext=ext)
