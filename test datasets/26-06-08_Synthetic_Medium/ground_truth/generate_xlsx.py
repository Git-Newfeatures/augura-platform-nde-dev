#!/usr/bin/env python3
"""
generate_xlsx.py
Reads ground_truth.json and produces a human-readable 4-sheet XLSX report.

Sheets:
  1. Summary       — dataset overview, severity counts, category histogram
  2. Findings      — one row per finding, full detail
  3. Cross-Refs    — every cross-table reference, flat
  4. Affected Rows — every (finding, row_id) pair, for grep-style lookups
"""

import json
import re
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import (
        PatternFill, Font, Alignment, Border, Side
    )
    from openpyxl.utils import get_column_letter
except ImportError:
    raise SystemExit("openpyxl not found — run: pip install openpyxl")

GT_DIR = Path(__file__).parent
SRC    = GT_DIR / "ground_truth.json"
OUT    = GT_DIR / "ground_truth_report.xlsx"


# ─── colours ────────────────────────────────────────────────────────────────
C_HARD    = "FFC7CE"   # red tint
C_SOFT    = "FFEB9C"   # amber tint
C_HEADER  = "1F4E79"   # dark blue
C_HEADER2 = "2E75B6"   # mid blue
C_ZEBRA   = "EBF3FB"   # light blue stripe
C_WHITE   = "FFFFFF"

def fill(hex_str):
    return PatternFill("solid", fgColor=hex_str)

def hdr_font(white=True):
    return Font(bold=True, color="FFFFFF" if white else "000000", size=10)

def wrap_align():
    return Alignment(wrap_text=True, vertical="top")

def thin_border():
    s = Side(style="thin", color="AAAAAA")
    return Border(left=s, right=s, top=s, bottom=s)

def set_col_width(ws, col_idx, width):
    ws.column_dimensions[get_column_letter(col_idx)].width = width


# ─── helpers ────────────────────────────────────────────────────────────────

def write_row(ws, row_idx, values, bg=None, bold=False, wrap=True):
    for col_idx, val in enumerate(values, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=str(val) if val is not None else "")
        cell.alignment = wrap_align() if wrap else Alignment(vertical="top")
        cell.border = thin_border()
        if bg:
            cell.fill = fill(bg)
        if bold:
            cell.font = Font(bold=True, size=10)


def write_header(ws, row_idx, labels, bg=C_HEADER):
    for col_idx, label in enumerate(labels, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=label)
        cell.fill = fill(bg)
        cell.font = hdr_font()
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border()
    ws.row_dimensions[row_idx].height = 30


def add_title(ws, text):
    ws.merge_cells(f"A1:{get_column_letter(20)}1")
    cell = ws["A1"]
    cell.value = text
    cell.font = Font(bold=True, size=14, color="FFFFFF")
    cell.fill = fill(C_HEADER)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28


# ════════════════════════════════════════════════════════════════════════════
# LOAD DATA
# ════════════════════════════════════════════════════════════════════════════

with open(SRC, encoding="utf-8") as f:
    gt = json.load(f)

meta     = gt["bundle_meta"]
findings = gt["findings"]


# ════════════════════════════════════════════════════════════════════════════
# WORKBOOK
# ════════════════════════════════════════════════════════════════════════════

wb = openpyxl.Workbook()
wb.remove(wb.active)   # remove default empty sheet


# ════════════════════════════════════════════════════════════════════════════
# SHEET 1: Summary
# ════════════════════════════════════════════════════════════════════════════

ws1 = wb.create_sheet("Summary")
add_title(ws1, f"Augura DQ Pressure-Test v2 — Ground Truth Summary")

row = 3

# Dataset info block
info_rows = [
    ("Dataset version",    meta["dataset_version"]),
    ("Patients",           meta["patients"]),
    ("Sites",              ", ".join(meta["sites"])),
    ("Timepoints",         ", ".join(meta["timepoints"])),
    ("Tables",             str(len(meta["tables"]))),
    ("Total findings",     meta["total_findings"]),
    ("Hard findings",      meta["severity_counts"]["hard"]),
    ("Soft findings",      meta["severity_counts"]["soft"]),
    ("Clean directory",    meta["clean_directory"]),
    ("Degraded directory", meta["degraded_directory"]),
]
write_header(ws1, row, ["Property", "Value"], bg=C_HEADER2)
row += 1
for i, (k, v) in enumerate(info_rows):
    bg = C_ZEBRA if i % 2 == 0 else C_WHITE
    write_row(ws1, row, [k, v], bg=bg)
    row += 1

row += 1

# Severity definitions
write_header(ws1, row, ["Severity", "Definition"], bg=C_HEADER2)
row += 1
for sev, defn in meta["severity_definitions"].items():
    bg = C_HARD if sev == "hard" else C_SOFT
    write_row(ws1, row, [sev.upper(), defn], bg=bg)
    row += 1

row += 1

# Category counts
from collections import Counter
cat_counts = Counter()
for f in findings:
    cat_counts[f["category"]] += 1

write_header(ws1, row, ["Category", "Count", "Severity"], bg=C_HEADER2)
row += 1
for i, (cat, cnt) in enumerate(sorted(cat_counts.items(), key=lambda x: -x[1])):
    # Find severity for this category
    sevs = list({f["severity"] for f in findings if f["category"] == cat})
    sev_str = "/".join(sorted(sevs))
    bg = C_HARD if sevs == ["hard"] else (C_SOFT if sevs == ["soft"] else C_ZEBRA)
    bg = C_ZEBRA if i % 2 == 0 else C_WHITE
    write_row(ws1, row, [cat, cnt, sev_str], bg=bg)
    row += 1

row += 1

# Tables affected
tables_affected = sorted({f["table"] for f in findings})
write_header(ws1, row, ["Tables with injected errors", "Findings count"], bg=C_HEADER2)
row += 1
for i, tbl in enumerate(tables_affected):
    cnt = sum(1 for f in findings if f["table"] == tbl)
    bg = C_ZEBRA if i % 2 == 0 else C_WHITE
    write_row(ws1, row, [tbl, cnt], bg=bg)
    row += 1

set_col_width(ws1, 1, 40)
set_col_width(ws1, 2, 60)
set_col_width(ws1, 3, 20)


# ════════════════════════════════════════════════════════════════════════════
# SHEET 2: Findings
# ════════════════════════════════════════════════════════════════════════════

ws2 = wb.create_sheet("Findings")
add_title(ws2, "Augura DQ Pressure-Test v2 — All Findings")

HEADERS = [
    "ID", "Severity", "Table", "Affected Row IDs",
    "Columns Involved", "Category",
    "What a Human Expert Would Observe",
    "Expected State",
    "Notes"
]
write_header(ws2, 3, HEADERS)
ws2.row_dimensions[3].height = 35

for i, f in enumerate(findings):
    row_idx = 4 + i
    sev = f["severity"]
    bg = C_HARD if sev == "hard" else C_SOFT

    primary_ids = ", ".join(f["affected_rows"]["primary_ids"])
    ref_ids     = f["affected_rows"].get("reference_ids", [])
    affected    = primary_ids
    if ref_ids:
        affected += f" (ref: {', '.join(ref_ids)})"

    cols = ", ".join(f["columns_involved"])

    values = [
        f["id"],
        sev.upper(),
        f["table"],
        affected,
        cols,
        f["category"],
        f["what_a_human_expert_would_observe"],
        f["expected_state"],
        f.get("notes", ""),
    ]
    write_row(ws2, row_idx, values, bg=bg)
    ws2.row_dimensions[row_idx].height = 90

col_widths = [6, 8, 28, 32, 32, 30, 70, 55, 40]
for ci, w in enumerate(col_widths, 1):
    set_col_width(ws2, ci, w)


# ════════════════════════════════════════════════════════════════════════════
# SHEET 3: Cross-References
# ════════════════════════════════════════════════════════════════════════════

ws3 = wb.create_sheet("Cross-References")
add_title(ws3, "Augura DQ Pressure-Test v2 — Cross-Table References")

XREF_HEADERS = [
    "Finding ID", "Primary Table", "Reference Table",
    "Reference Row ID", "Reference Column", "Context"
]
write_header(ws3, 3, XREF_HEADERS)

xrow = 4
for f in findings:
    for xr in f.get("cross_references", []):
        bg = C_ZEBRA if xrow % 2 == 0 else C_WHITE
        write_row(ws3, xrow, [
            f["id"],
            f["table"],
            xr["table"],
            xr["row_id"],
            xr["column"],
            xr["context"],
        ], bg=bg)
        xrow += 1

if xrow == 4:
    ws3.cell(row=4, column=1, value="(no cross-references recorded)")

x_widths = [10, 28, 28, 22, 28, 60]
for ci, w in enumerate(x_widths, 1):
    set_col_width(ws3, ci, w)


# ════════════════════════════════════════════════════════════════════════════
# SHEET 4: Affected Rows
# ════════════════════════════════════════════════════════════════════════════

ws4 = wb.create_sheet("Affected Rows")
add_title(ws4, "Augura DQ Pressure-Test v2 — Affected Row Index")

AR_HEADERS = [
    "Finding ID", "Severity", "Table", "Row ID", "ID Column",
    "Role", "Category", "Short description"
]
write_header(ws4, 3, AR_HEADERS)

arow = 4
for f in findings:
    ar = f["affected_rows"]
    sev = f["severity"]
    short = f["what_a_human_expert_would_observe"][:120].replace("\n", " ")

    for rid in ar["primary_ids"]:
        bg = C_HARD if sev == "hard" else C_SOFT
        write_row(ws4, arow, [
            f["id"], sev.upper(), f["table"], rid, ar["id_column"],
            "primary", f["category"], short
        ], bg=bg)
        arow += 1

    for rid in ar.get("reference_ids", []):
        write_row(ws4, arow, [
            f["id"], sev.upper(), f["table"], rid, ar["id_column"],
            "reference", f["category"], short
        ], bg=C_ZEBRA)
        arow += 1

ar_widths = [10, 8, 28, 14, 14, 10, 30, 80]
for ci, w in enumerate(ar_widths, 1):
    set_col_width(ws4, ci, w)


# ════════════════════════════════════════════════════════════════════════════
# FREEZE PANES & SAVE
# ════════════════════════════════════════════════════════════════════════════

for ws in [ws2, ws3, ws4]:
    ws.freeze_panes = ws.cell(row=4, column=1)

wb.save(OUT)
print(f"✅  Saved: {OUT}")
print(f"    Sheets: Summary, Findings ({len(findings)} rows), Cross-References, Affected Rows")
